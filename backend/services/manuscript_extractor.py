"""
GapGuard AI — Manuscript Text Extraction Service

Extracts readable text from research manuscripts (.pdf, .docx, .txt).
Preserves paragraph boundaries, normalizes whitespace, and rejects corrupted or empty files.
Does NOT perform AI reasoning, gap detection, or plagiarism checks.
"""

import io
import re
import zipfile
from typing import Dict, Any

import pdfplumber
from docx import Document
from docx.opc.exceptions import PackageNotFoundError


MAX_MANUSCRIPT_SIZE_BYTES = 15 * 1024 * 1024  # 15 MB
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def normalize_whitespace(text: str) -> str:
    """
    Normalizes extracted manuscript text:
    - Normalizes Windows and legacy Mac line endings to standard Unix newlines.
    - Strips trailing whitespace per line.
    - Collapses runs of more than two consecutive newlines into double newlines (preserving paragraphs).
    - Strips leading and trailing overall whitespace.
    """
    if not text:
        return ""

    # Normalize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Strip trailing whitespace on each line
    lines = [line.rstrip() for line in text.split("\n")]
    text = "\n".join(lines)

    # Collapse multiple blank lines (> 2 newlines into 2)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def extract_text_from_txt(file_bytes: bytes) -> str:
    """
    Extract readable text from a plain-text manuscript document.
    Tries utf-8, utf-8-sig, and latin-1 encodings.
    """
    if not file_bytes or len(file_bytes.strip()) == 0:
        raise ValueError("The TXT manuscript document is empty.")

    decoded = None
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            decoded = file_bytes.decode(encoding)
            break
        except (UnicodeDecodeError, LookupError):
            continue

    if decoded is None:
        decoded = file_bytes.decode("utf-8", errors="replace")

    clean_text = normalize_whitespace(decoded)
    if not clean_text:
        raise ValueError("The TXT manuscript document contains no readable text.")

    return clean_text


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract text from a PDF manuscript using pdfplumber.
    Preserves page flow and paragraph boundaries.
    """
    if not file_bytes or len(file_bytes) == 0:
        raise ValueError("The PDF manuscript document is empty (0 bytes).")

    # Basic magic byte check for PDF (%PDF-)
    if not file_bytes.startswith(b"%PDF"):
        raise ValueError("The file does not appear to be a valid PDF document (missing PDF header).")

    pages_text = []

    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            if not pdf.pages:
                raise ValueError("The PDF manuscript document contains no pages.")

            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    pages_text.append(page_text.strip())

    except Exception as exc:
        if isinstance(exc, ValueError):
            raise
        raise ValueError(
            f"The PDF manuscript document is corrupted or unreadable: {str(exc)}"
        ) from exc

    if not pages_text:
        raise ValueError(
            "The PDF manuscript contains no readable text. Scanned image-only PDFs without OCR are not supported in this phase."
        )

    full_text = "\n\n".join(pages_text)
    clean_text = normalize_whitespace(full_text)

    if not clean_text:
        raise ValueError("The PDF manuscript document contains no readable text.")

    return clean_text


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract text from a DOCX manuscript using python-docx.
    Preserves paragraph structure and extracts table contents.
    """
    if not file_bytes or len(file_bytes) == 0:
        raise ValueError("The DOCX manuscript document is empty (0 bytes).")

    try:
        document = Document(io.BytesIO(file_bytes))
    except (PackageNotFoundError, zipfile.BadZipFile, Exception) as exc:
        raise ValueError(
            f"The DOCX manuscript document is corrupted or unreadable: {str(exc)}"
        ) from exc

    elements = []

    # Extract paragraphs
    for paragraph in document.paragraphs:
        p_text = paragraph.text.strip()
        if p_text:
            elements.append(p_text)

    # Extract tables if present
    for table in document.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Deduplicate identical adjacent cells caused by merged cells
                deduped = []
                for c in row_cells:
                    if not deduped or deduped[-1] != c:
                        deduped.append(c)
                if deduped:
                    elements.append(" | ".join(deduped))

    if not elements:
        raise ValueError("The DOCX manuscript document is empty and contains no readable text.")

    full_text = "\n\n".join(elements)
    clean_text = normalize_whitespace(full_text)

    if not clean_text:
        raise ValueError("The DOCX manuscript document contains no readable text.")

    return clean_text


def extract_manuscript_text(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Main dispatch function for manuscript text extraction.
    Validates file extension, size, and dispatches to format-specific extractors.
    Returns structured metadata including character and word counts.
    """
    if not filename or "." not in filename:
        raise ValueError("Manuscript filename must include an extension (.pdf, .docx, or .txt).")

    ext = "." + filename.rsplit(".", 1)[-1].lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file format '{ext}'. Only PDF (.pdf), DOCX (.docx), and TXT (.txt) manuscripts are supported."
        )

    if not file_bytes:
        raise ValueError(f"The uploaded {ext.upper().replace('.', '')} manuscript file is empty (0 bytes).")

    if len(file_bytes) > MAX_MANUSCRIPT_SIZE_BYTES:
        raise ValueError(
            f"Manuscript file exceeds the maximum allowed size of {MAX_MANUSCRIPT_SIZE_BYTES // (1024 * 1024)} MB."
        )

    file_type = "PDF" if ext == ".pdf" else "DOCX" if ext == ".docx" else "TXT"

    if ext == ".pdf":
        clean_text = extract_text_from_pdf(file_bytes)
    elif ext == ".docx":
        clean_text = extract_text_from_docx(file_bytes)
    elif ext == ".txt":
        clean_text = extract_text_from_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported format '{ext}'.")

    # Word and character count computation
    character_count = len(clean_text)
    word_count = len(clean_text.split())

    return {
        "success": True,
        "file_name": filename,
        "file_type": file_type,
        "character_count": character_count,
        "word_count": word_count,
        "text": clean_text,
    }
