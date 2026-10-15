import io
import os
import sys
import uuid
import hashlib
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document

# Ensure backend directory is in sys.path
BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import database.database as db_module
import routes.reports as reports_routes
import routes.courses as courses_routes
import routes.student as student_routes
import routes.submissions as submissions_routes
import routes.auth as auth_routes
import routes.custom_references as custom_refs_routes

from main import app
from fastapi.testclient import TestClient
from database.database import initialize_database
from models.user import create_user
from models.course import create_course, enroll_student
from utils.auth import create_access_token


class TestCustomReferencesFeature(unittest.TestCase):
    """
    Comprehensive isolated test suite for custom reference resources.
    Tests custom reference CRUD, input validation, file uploads (TXT, DOCX),
    authorization isolation, and plagiarism engine integration.
    Guarantees test safety by using a temporary isolated SQLite database.
    """

    @classmethod
    def setUpClass(cls):
        # 1. Record live database hash for safety check
        cls.live_db_path = BACKEND_DIR / "academic_integrity.db"
        if cls.live_db_path.exists():
            cls.live_db_initial_hash = hashlib.sha256(cls.live_db_path.read_bytes()).hexdigest()
        else:
            cls.live_db_initial_hash = None

        # 2. Create isolated temporary directory & SQLite database
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.temp_db_path = Path(cls.temp_dir.name) / "test_custom_refs.db"

        def get_isolated_connection():
            conn = sqlite3.connect(cls.temp_db_path)
            conn.execute("PRAGMA foreign_keys = ON;")
            conn.row_factory = sqlite3.Row
            return conn

        cls.patchers = [
            patch.object(db_module, "DATABASE_PATH", cls.temp_db_path),
            patch.object(db_module, "get_connection", side_effect=get_isolated_connection),
            patch.object(reports_routes, "get_connection", side_effect=get_isolated_connection),
            patch.object(courses_routes, "get_connection", side_effect=get_isolated_connection),
            patch.object(student_routes, "get_connection", side_effect=get_isolated_connection),
            patch.object(submissions_routes, "get_connection", side_effect=get_isolated_connection),
            patch.object(auth_routes, "get_connection", side_effect=get_isolated_connection),
            patch.object(custom_refs_routes, "get_connection", side_effect=get_isolated_connection),
        ]
        for p in cls.patchers:
            p.start()

        # Initialize schema in isolated DB
        initialize_database()

        cls.client = TestClient(app)

        # Create isolated users & course
        conn = get_isolated_connection()
        try:
            cls.scholar_a_id = create_user(conn, "Scholar Alice", "alice.refs@example.com", "password123", "student")
            cls.scholar_b_id = create_user(conn, "Scholar Bob", "bob.refs@example.com", "password123", "student")
            cls.prof_id = create_user(conn, "Prof Xavier", "xavier@example.com", "password123", "professor")
            cls.course_id = create_course(conn, "Natural Language Processing", "CS802", "NLP area", cls.prof_id)
            enroll_student(conn, cls.course_id, cls.scholar_a_id)
            enroll_student(conn, cls.course_id, cls.scholar_b_id)
        finally:
            conn.close()

        cls.token_a = create_access_token({
            "sub": str(cls.scholar_a_id),
            "email": "alice.refs@example.com",
            "role": "student"
        })
        cls.token_b = create_access_token({
            "sub": str(cls.scholar_b_id),
            "email": "bob.refs@example.com",
            "role": "student"
        })

    @classmethod
    def tearDownClass(cls):
        for p in cls.patchers:
            p.stop()
        cls.temp_dir.cleanup()

        # Verify live database was never touched
        if cls.live_db_initial_hash and cls.live_db_path.exists():
            final_hash = hashlib.sha256(cls.live_db_path.read_bytes()).hexdigest()
            assert cls.live_db_initial_hash == final_hash, "CRITICAL: Live database was modified during testing!"

    def test_01_add_custom_reference_via_pasted_text(self):
        """1. Scholar can add custom reference paper through pasted text."""
        payload = {
            "title": "Attention and Multi-Head Cross Attention Foundations",
            "author": "Ashish Vaswani, Noam Shazeer",
            "publication_year": "2017",
            "research_field": "Natural Language Processing",
            "text_content": (
                "The dominant sequence transduction models are based on complex recurrent or convolutional neural networks. "
                "We propose the Transformer, a model architecture eschewing recurrence and relying entirely on an attention mechanism "
                "to draw global dependencies between input and output representations."
            )
        }
        res = self.client.post(
            "/api/references/custom",
            headers={"Authorization": f"Bearer {self.token_a}"},
            data=payload
        )
        self.assertEqual(res.status_code, 201)
        data = res.json()
        self.assertTrue(data.get("success"))
        self.assertEqual(data["reference"]["title"], payload["title"])
        self.assertEqual(data["reference"]["resource_type"], "My References")
        self.assertIn("dominant sequence transduction", data["reference"]["preview"])

    def test_02_add_custom_reference_via_docx_file(self):
        """2. Scholar can upload DOCX reference file and extract text."""
        docx_io = io.BytesIO()
        doc = Document()
        doc.add_paragraph("Diffusion models generate high fidelity images by reversing a forward Gaussian noise process.")
        doc.add_paragraph("Classifier-free guidance significantly improves sample quality without requiring separate classifier training.")
        doc.save(docx_io)
        docx_bytes = docx_io.getvalue()

        res = self.client.post(
            "/api/references/custom",
            headers={"Authorization": f"Bearer {self.token_a}"},
            data={
                "title": "Denoising Diffusion Probabilistic Models Analysis",
                "author": "Jonathan Ho et al.",
                "publication_year": "2020",
                "research_field": "Computer Vision",
            },
            files={"file": ("diffusion_paper.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        self.assertEqual(res.status_code, 201)
        data = res.json()
        self.assertTrue(data.get("success"))
        self.assertEqual(data["reference"]["filename"], "diffusion_paper.docx")

    def test_03_add_custom_reference_via_txt_file(self):
        """3. Scholar can upload TXT reference file and extract text."""
        txt_content = (
            "Contrastive Language-Image Pre-Training models transfer zero-shot knowledge to downstream classification tasks. "
            "Natural language supervision enables learning visual concepts at scale directly from web data."
        ).encode("utf-8")

        res = self.client.post(
            "/api/references/custom",
            headers={"Authorization": f"Bearer {self.token_a}"},
            data={"title": "Learning Transferable Visual Models From Natural Language"},
            files={"file": ("clip_paper.txt", txt_content, "text/plain")}
        )
        self.assertEqual(res.status_code, 201)
        self.assertTrue(res.json().get("success"))

    def test_04_validation_missing_title_or_content(self):
        """4. Validation properly rejects empty title, unsupported files, and empty text."""
        # Missing title
        res1 = self.client.post(
            "/api/references/custom",
            headers={"Authorization": f"Bearer {self.token_a}"},
            data={"title": "   ", "text_content": "Some valid content."}
        )
        self.assertEqual(res1.status_code, 400)
        self.assertIn("title is required", res1.json()["detail"].lower())

        # Missing both file and text
        res2 = self.client.post(
            "/api/references/custom",
            headers={"Authorization": f"Bearer {self.token_a}"},
            data={"title": "Empty Content Paper"}
        )
        self.assertEqual(res2.status_code, 400)

        # Unsupported file format
        res3 = self.client.post(
            "/api/references/custom",
            headers={"Authorization": f"Bearer {self.token_a}"},
            data={"title": "Binary Executable"},
            files={"file": ("test.bin", b"\x00\x01\x02", "application/octet-stream")}
        )
        self.assertEqual(res3.status_code, 400)
        self.assertIn("unsupported", res3.json()["detail"].lower())

    def test_05_list_custom_references_and_scholar_isolation(self):
        """5. Scholar A can view their references, and Scholar B does not see Scholar A's references."""
        # Scholar A lists their references
        res_a = self.client.get(
            "/api/references/custom",
            headers={"Authorization": f"Bearer {self.token_a}"}
        )
        self.assertEqual(res_a.status_code, 200)
        refs_a = res_a.json()["references"]
        self.assertGreaterEqual(len(refs_a), 3)

        # Scholar B lists their references (should be 0)
        res_b = self.client.get(
            "/api/references/custom",
            headers={"Authorization": f"Bearer {self.token_b}"}
        )
        self.assertEqual(res_b.status_code, 200)
        self.assertEqual(len(res_b.json()["references"]), 0)

    def test_06_unauthorized_deletion_prevented(self):
        """6. Scholar B cannot delete Scholar A's private reference."""
        # Get one of Scholar A's reference IDs
        res_a = self.client.get("/api/references/custom", headers={"Authorization": f"Bearer {self.token_a}"})
        ref_id = res_a.json()["references"][0]["id"]

        # Scholar B attempts deletion
        res_delete_b = self.client.delete(
            f"/api/references/custom/{ref_id}",
            headers={"Authorization": f"Bearer {self.token_b}"}
        )
        self.assertEqual(res_delete_b.status_code, 403)

    def test_07_custom_reference_included_in_plagiarism_check(self):
        """7. Plagiarism detection compares uploaded manuscript against scholar's custom reference."""
        # Scholar A uploads manuscript containing exact text from their added diffusion reference
        matching_manuscript = (
            "In this research chapter, we explore modern generative techniques. "
            "Diffusion models generate high fidelity images by reversing a forward Gaussian noise process. "
            "Classifier-free guidance significantly improves sample quality without requiring separate classifier training. "
            "Our empirical evaluation examines sample diversity across diverse multimodal benchmarks."
        ).encode("utf-8")

        res_upload = self.client.post(
            "/api/submissions",
            headers={"Authorization": f"Bearer {self.token_a}"},
            data={"title": "Generative Diffusion Study", "course_id": str(self.course_id)},
            files={"file": ("manuscript_diffusion.txt", matching_manuscript, "text/plain")}
        )
        self.assertEqual(res_upload.status_code, 201)
        data = res_upload.json()

        # Verify custom reference matched
        matches = data["matches"]
        custom_matches = [m for m in matches if m.get("type") == "My Reference Resource"]
        self.assertGreater(len(custom_matches), 0, "Custom reference should be in comparison matches")

        top_custom = custom_matches[0]
        self.assertIn("Denoising Diffusion", top_custom["title"])
        self.assertGreater(top_custom["similarity_percentage"], 30.0)
        self.assertGreater(len(top_custom["matched_segments"]), 0)

        # Check evidence snippets
        seg = top_custom["matched_segments"][0]
        self.assertIn("Diffusion models generate high fidelity images", seg["target_snippet"])

    def test_08_scholar_b_does_not_match_scholar_a_private_reference(self):
        """8. Scholar B's plagiarism checks do not include Scholar A's private custom references."""
        matching_manuscript = (
            "Diffusion models generate high fidelity images by reversing a forward Gaussian noise process. "
            "Classifier-free guidance significantly improves sample quality without requiring separate classifier training."
        ).encode("utf-8")

        res_upload_b = self.client.post(
            "/api/submissions",
            headers={"Authorization": f"Bearer {self.token_b}"},
            data={"title": "Bob's Independent Diffusion Study", "course_id": str(self.course_id)},
            files={"file": ("bob_manuscript.txt", matching_manuscript, "text/plain")}
        )
        self.assertEqual(res_upload_b.status_code, 201)
        data_b = res_upload_b.json()

        # Scholar B has no custom references, so none should appear in matches with type "My Reference Resource"
        custom_matches_b = [m for m in data_b["matches"] if m.get("type") == "My Reference Resource"]
        self.assertEqual(len(custom_matches_b), 0)

    def test_09_authorized_deletion_succeeds(self):
        """9. Scholar A can delete their own custom reference."""
        # Add a temporary reference to delete
        res_add = self.client.post(
            "/api/references/custom",
            headers={"Authorization": f"Bearer {self.token_a}"},
            data={"title": "Temporary Paper to Delete", "text_content": "Temporary text for deletion test."}
        )
        ref_id = res_add.json()["reference"]["id"]

        # Delete it
        res_del = self.client.delete(
            f"/api/references/custom/{ref_id}",
            headers={"Authorization": f"Bearer {self.token_a}"}
        )
        self.assertEqual(res_del.status_code, 200)
        self.assertTrue(res_del.json().get("success"))

        # Confirm it's gone
        res_del_again = self.client.delete(
            f"/api/references/custom/{ref_id}",
            headers={"Authorization": f"Bearer {self.token_a}"}
        )
        self.assertEqual(res_del_again.status_code, 404)

    def test_10_upload_multiple_reference_documents_in_analysis(self):
        """10. Scholar can upload paper with multiple reference documents, comparing against only those references."""
        paper_text = (
            "This study investigates neural representation alignment across vision-language architectures. "
            "Vision transformers tokenize images into non-overlapping patch embeddings projected into unified linear vectors. "
            "Contrastive learning aligns visual tokens with textual embeddings in a joint representation space."
        ).encode("utf-8")

        ref1_text = (
            "Vision transformers tokenize images into non-overlapping patch embeddings projected into unified linear vectors."
        ).encode("utf-8")

        ref2_text = (
            "Contrastive learning aligns visual tokens with textual embeddings in a joint representation space."
        ).encode("utf-8")

        files = [
            ("file", ("multimodal_paper.txt", paper_text, "text/plain")),
            ("reference_files", ("vit_ref.txt", ref1_text, "text/plain")),
            ("reference_files", ("clip_ref.txt", ref2_text, "text/plain")),
        ]

        res = self.client.post(
            "/api/submissions",
            headers={"Authorization": f"Bearer {self.token_b}"},
            data={"title": "Multimodal Alignment Study", "course_id": str(self.course_id)},
            files=files
        )
        self.assertEqual(res.status_code, 201)
        data = res.json()
        self.assertTrue(data.get("success"))

        matches = data["matches"]
        matched_titles = [m.get("title", "") for m in matches]

        # Both uploaded references must be present in comparison
        self.assertTrue(any("vit_ref" in t for t in matched_titles), "vit_ref should be in matches")
        self.assertTrue(any("clip_ref" in t for t in matched_titles), "clip_ref should be in matches")

        # Crucially: no hardcoded system references should be in matches
        self.assertFalse(any("P001" in m.get("reference_id", "") or "P002" in m.get("reference_id", "") for m in matches))

        # Check evidence segments matched
        evidence_summary = data.get("evidence_details", {}).get("similarity_evidence", {})
        self.assertGreaterEqual(evidence_summary.get("reference_sources_compared", 0), 2)

    def test_11_no_hardcoded_system_references_in_plagiarism_check(self):
        """11. Plagiarism check against zero references yields 0% similarity and 0 reference sources compared."""
        # Create a new isolated student with no custom references
        conn = sqlite3.connect(self.temp_db_path)
        conn.row_factory = sqlite3.Row
        try:
            scholar_c_id = create_user(conn, "Scholar Charlie", "charlie@example.com", "password123", "student")
            empty_course_id = create_course(conn, "Quantum Computing", "QC101", "Quantum", self.prof_id)
            enroll_student(conn, empty_course_id, scholar_c_id)
        finally:
            conn.close()

        token_c = create_access_token({
            "sub": str(scholar_c_id),
            "email": "charlie@example.com",
            "role": "student"
        })

        unique_paper = (
            "Quantum thermodynamics examines how microscopic systems obey quantum laws under macroscopic thermal constraints. "
            "We analyze entropy production in non-equilibrium spin chains using isolated Hamiltonian evolution."
        ).encode("utf-8")

        res = self.client.post(
            "/api/submissions",
            headers={"Authorization": f"Bearer {token_c}"},
            data={"title": "Quantum Thermodynamics Notes", "course_id": str(empty_course_id)},
            files={"file": ("quantum_notes.txt", unique_paper, "text/plain")}
        )
        self.assertEqual(res.status_code, 201)
        data = res.json()

        # Overall similarity should be 0.0% and no references compared
        self.assertEqual(data["overall_similarity_score"], 0.0)
        self.assertEqual(len(data["matches"]), 0)
        evidence_summary = data.get("evidence_details", {}).get("similarity_evidence", {})
        self.assertEqual(evidence_summary.get("reference_sources_compared", 0), 0)
        self.assertEqual(evidence_summary.get("sources_compared", 0), 0)
        self.assertFalse(any("P0" in m.get("reference_id", "") for m in data["matches"]))


if __name__ == "__main__":
    unittest.main()

